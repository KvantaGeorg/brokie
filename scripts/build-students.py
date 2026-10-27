"""Convert the Brokie student dataset (dataset.xlsx) into src/data/students/students.json for the demo app.

Usage: python3 scripts/build-students.py path/to/dataset.xlsx

The sheet has no transaction dates. Each student has three months of rows in order; a month starts at the
first of that month's recurring payments (rent, utilities, phone plan, transport pass, parents' support).
Rows are then spread evenly over a 30-day month, so "day" is an approximation used for weekly views.
"""
import json
import sys
from pathlib import Path

import openpyxl

DAYS = 30
RECURRING = {"Rent", "Utilities", "Phone plan", "Transport pass", "Parents' support"}
OUT = Path(__file__).resolve().parent.parent / "src" / "data" / "students" / "students.json"


def month_starts(rows):
    starts = []
    for k in range(3):
        candidates = []
        for cat in RECURRING:
            idx = [i for i, r in enumerate(rows) if r["category"] == cat]
            if len(idx) == 3:
                candidates.append(idx[k])
        starts.append(min(candidates) if candidates else round(k * len(rows) / 3))
    starts[0] = 0
    return starts


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]

    raw_cats = sorted({r["category"] for r in rows})
    by_user = {}
    for r in rows:
        by_user.setdefault(r["user_id"], []).append(r)

    students = []
    for uid, urows in by_user.items():
        urows.sort(key=lambda r: r["tx_id"])
        starts = month_starts(urows)
        bounds = starts[1:] + [len(urows)]
        first = urows[0]
        subs = [s.strip() for s in (first["subscriptions"] or "").split(",") if s.strip()]

        tx = []
        budgets = {}
        for m in range(3):
            month_rows = urows[starts[m]:bounds[m]]
            n = len(month_rows)
            sub_i = 0
            month_subs = [r for r in month_rows if r["category"] == "Subscriptions"]
            for i, r in enumerate(month_rows):
                day = min(DAYS, i * DAYS // n + 1)
                item = [m, day, raw_cats.index(r["category"]), round(r["amount_eur"] * 100)]
                if r["category"] == "Subscriptions" and len(month_subs) == len(subs):
                    item.append(subs[sub_i])
                    sub_i += 1
                tx.append(item)
                if r["category_budget_eur"] is not None:
                    budgets[r["category"]] = r["category_budget_eur"]

        students.append(
            {
                "id": uid,
                "handle": first["username"],
                "age": first["age"],
                "city": first["city"],
                "university": first["university"],
                "studyYear": first["study_year"],
                "living": first["living_situation"],
                "language": first["language"],
                "bank": first["bank_name"],
                "tone": first["notification_tone"].lower(),
                "frequency": first["notification_frequency"],
                "income": {
                    "parents": first["income_parents_support_eur"] or 0,
                    "job": first["income_part-time_job_eur"] or 0,
                    "scholarship": first["income_scholarship_eur"] or 0,
                    "other": first["income_other_eur"] or 0,
                    "total": round(first["income_total_eur"], 2),
                },
                "subscriptions": subs,
                "goal": {
                    "name": first["savings_goal"],
                    "target": first["goal_target_eur"],
                    "saved": first["goal_saved_so_far_eur"],
                    "date": first["goal_deadline"].date().isoformat() if first["goal_deadline"] else None,
                },
                "budgets": budgets,
                "tx": tx,
            }
        )

    OUT.write_text(json.dumps({"rawCategories": raw_cats, "days": DAYS, "students": students}, ensure_ascii=False, separators=(",", ":")))
    print(f"{len(students)} students, {len(rows)} transactions -> {OUT} ({OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "dataset.xlsx")
